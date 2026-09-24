import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

E = 914400
ORANGE = RGBColor(0xE8, 0x96, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1E, 0x3A, 0x5C)
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = RGBColor(0x5A, 0x5F, 0x66)


# ---------------------------------------------------------------- helpers
def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def shape_with_text(slide, fragment):
    for sh in slide.shapes:
        if sh.has_text_frame and fragment in sh.text_frame.text:
            return sh
    raise KeyError(fragment)


def set_title(slide, text):
    tf = slide.shapes.title.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def set_run_text(shape, text):
    """Replace a text box's text, keeping the first run's formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def replace_in_shape(shape, old, new):
    done = False
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                done = True
    if not done:
        raise KeyError(f"'{old}' not found in {shape.text_frame.text[:60]!r}")


def set_bullets(placeholder, items, size=16):
    """items: list of str or (str, level). Uses the placeholder's own bullet styles."""
    tf = placeholder.text_frame
    tf.clear()
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text
        p.level = lvl
        p.space_after = Pt(8)
        for r in p.runs:
            r.font.size = Pt(size if lvl == 0 else size - 2)


def delete_paragraph_containing(shape, fragment):
    for p in shape.text_frame.paragraphs:
        if fragment in "".join(r.text for r in p.runs):
            p._p.getparent().remove(p._p)
            return
    raise KeyError(fragment)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def add_label(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.92), Inches(0.14), Inches(6), Inches(0.32))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text.upper()
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = ORANGE
    rPr = r._r.get_or_add_rPr()
    rPr.set("spc", "150")
    tb.name = "Section label"


def set_notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        from lxml import etree
        P = "http://schemas.openxmlformats.org/presentationml/2006/main"
        A = "http://schemas.openxmlformats.org/drawingml/2006/main"
        xml = (f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="3" name="Notes Placeholder 2"/>'
               '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
               '<p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody></p:sp>')
        ns.shapes._spTree.append(etree.fromstring(xml))
        tf = ns.notes_text_frame
    tf.text = text


def new_content_slide(pr, layout, title, bullets, size=16):
    s = pr.slides.add_slide(layout)
    s.shapes.title.text = title
    body = [ph for ph in s.placeholders if ph.placeholder_format.idx != 0][0]
    set_bullets(body, bullets, size)
    return s


def new_title_only_slide(pr, layout, title):
    s = pr.slides.add_slide(layout)
    s.shapes.title.text = title
    for ph in [ph for ph in s.placeholders if ph.placeholder_format.idx != 0]:
        ph._element.getparent().remove(ph._element)
    return s


def add_picture_fit(slide, path, top=1.95, max_h=4.45, max_w=11.2):
    w, h = Image.open(path).size
    ratio = w / h
    width = min(max_w, max_h * ratio)
    height = width / ratio
    left = (13.333 - width) / 2
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))


def card(slide, x, y, w, h, lines, fill=WHITE, text_color=NAVY, border=None):
    """lines: list of (text, size, bold, color or None)."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    if border:
        box.line.color.rgb = border
        box.line.width = Pt(1.5)
    else:
        box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color or text_color
    return box


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = ORANGE
    c.line.width = Pt(2.25)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def flow_slide(pr, layout, stages, exclusions):
    """stages: list of (big number, caption); exclusions: list aligned between stages (text or None)."""
    s = new_title_only_slide(pr, layout, "Participant Flow")
    n = len(stages)
    top, gap_y = 1.95, 0.55
    box_h = (4.55 - gap_y * (n - 1)) / n
    x, w = 2.3, 4.2
    for i, (num, cap) in enumerate(stages):
        y = top + i * (box_h + gap_y)
        card(s, x, y, w, box_h, [(num, 34, True, RED if i == n - 1 else NAVY), (cap, 16, False, GREY)])
        if i < n - 1:
            y_mid = y + box_h + gap_y / 2
            arrow(s, x + w / 2, y + box_h, x + w / 2, y + box_h + gap_y)
            if exclusions[i]:
                c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w / 2), Inches(y_mid), Inches(7.2), Inches(y_mid))
                c.line.color.rgb = ORANGE
                c.line.width = Pt(1.5)
                card(s, 7.2, y_mid - 0.6, 4.8, 1.2, [(exclusions[i], 15, False, WHITE)],
                     fill=RGBColor(0x1B, 0x2A, 0x4A), border=ORANGE)
    return s


def reorder_and_prune(pr, order):
    lst = pr.slides._sldIdLst
    by_part = {}
    for sid in list(lst):
        by_part[pr.part.related_part(sid.rId)] = sid
    keep = [by_part[s.part] for s in order]
    for sid in list(lst):
        if sid not in keep:
            pr.part.drop_rel(sid.rId)
        lst.remove(sid)
    for sid in keep:
        lst.append(sid)


def label_all(order, sections):
    for idx, s in enumerate(order, start=1):
        for (lo, hi, name) in sections:
            if lo <= idx <= hi:
                add_label(s, name)
